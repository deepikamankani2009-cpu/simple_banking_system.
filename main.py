import os
from openpyxl import Workbook, load_workbook

EXCEL_FILE = "banking_system.xlsx"


def initialize_excel():
    """Excel file setup karta hai agar pehle se na bani ho."""
    if not os.path.exists(EXCEL_FILE):
        wb = Workbook()
        ws = wb.active
        ws.title = "Accounts"

        # Headers
        ws.append(["Account Number", "Name", "Mobile Number", "PIN", "Balance"])

        # Default Data (Project requirement ke according)
        ws.append([1001, "Rahul Patil", "9876543210", "1234", 5000])
        ws.append([1002, "Priya Sharma", "9876543211", "5678", 3000])

        wb.save(EXCEL_FILE)


def create_account():
    print("\n----------------------------------")
    print("        CREATE NEW ACCOUNT        ")
    print("----------------------------------")
    name = input("Enter Account Holder Name: ").strip()
    mobile = input("Enter 10-digit Mobile Number: ").strip()

    if not (mobile.isdigit() and len(mobile) == 10):
        print("Error: Invalid Mobile Number! Exactly 10 digits required.")
        return

    pin = input("Enter 4-digit PIN: ").strip()
    if not (pin.isdigit() and len(pin) == 4):
        print("Error: Invalid PIN! Exactly 4 digits required.")
        return

    try:
        initial_deposit = float(input("Enter Initial Deposit Amount: "))
        if initial_deposit < 0:
            print("Error: Deposit amount cannot be negative!")
            return
    except ValueError:
        print("Error: Invalid input for amount!")
        return

    wb = load_workbook(EXCEL_FILE)
    ws = wb.active

    # Auto generate next Account Number
    max_acc = 1000
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] and isinstance(row[0], int) and row[0] > max_acc:
            max_acc = row[0]

    new_acc_num = max_acc + 1

    # Excel me row add karna
    ws.append([new_acc_num, name, mobile, pin, initial_deposit])
    wb.save(EXCEL_FILE)

    print("\nAccount Created Successfully!")
    print(f"Generated Account Number: {new_acc_num}")


def deposit_money():
    print("\n----------------------------------")
    print("          DEPOSIT MONEY           ")
    print("----------------------------------")
    try:
        acc_num = int(input("Enter Account Number: "))
    except ValueError:
        print("Error: Account number must be numeric!")
        return

    wb = load_workbook(EXCEL_FILE)
    ws = wb.active

    found_row = None
    account_data = None
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row[0] == acc_num:
            found_row = row_idx
            account_data = row
            break

    if not account_data:
        print("Error: Account Not Found!")
        return

    pin = input("Enter 4-digit PIN: ").strip()
    if str(pin) != str(account_data[3]):
        print("Error: Incorrect PIN!")
        return

    try:
        amount = float(input("Enter Deposit Amount: "))
        if amount <= 0:
            print("Error: Deposit amount must be positive!")
            return
    except ValueError:
        print("Error: Invalid amount format!")
        return

    current_bal = float(account_data[4])
    new_bal = current_bal + amount
    ws.cell(row=found_row, column=5, value=new_bal)
    wb.save(EXCEL_FILE)

    print("\nDeposit Successful!")
    print(f"New Balance: Rs. {new_bal:.2f}")


def withdraw_money():
    print("\n----------------------------------")
    print("         WITHDRAW MONEY           ")
    print("----------------------------------")
    try:
        acc_num = int(input("Enter Account Number: "))
    except ValueError:
        print("Error: Account number must be numeric!")
        return

    wb = load_workbook(EXCEL_FILE)
    ws = wb.active

    found_row = None
    account_data = None
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row[0] == acc_num:
            found_row = row_idx
            account_data = row
            break

    if not account_data:
        print("Error: Account Not Found!")
        return

    pin = input("Enter 4-digit PIN: ").strip()
    if str(pin) != str(account_data[3]):
        print("Error: Incorrect PIN!")
        return

    current_bal = float(account_data[4])
    try:
        amount = float(input("Enter Withdrawal Amount: "))
        if amount <= 0:
            print("Error: Amount must be positive!")
            return
    except ValueError:
        print("Error: Invalid amount format!")
        return

    if amount > current_bal:
        print(f"Error: Insufficient Balance! Available Balance: Rs. {current_bal:.2f}")
        return

    new_bal = current_bal - amount
    ws.cell(row=found_row, column=5, value=new_bal)
    wb.save(EXCEL_FILE)

    print("\nWithdrawal Successful!")
    print(f"Remaining Balance: Rs. {new_bal:.2f}")


def check_balance():
    print("\n----------------------------------")
    print("          CHECK BALANCE           ")
    print("----------------------------------")
    try:
        acc_num = int(input("Enter Account Number: "))
    except ValueError:
        print("Error: Account number must be numeric!")
        return

    wb = load_workbook(EXCEL_FILE)
    ws = wb.active

    account_data = None
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] == acc_num:
            account_data = row
            break

    if not account_data:
        print("Error: Account Not Found!")
        return

    pin = input("Enter 4-digit PIN: ").strip()
    if str(pin) != str(account_data[3]):
        print("Error: Incorrect PIN!")
        return

    print("\n======== ACCOUNT DETAILS ========")
    print(f"Account Number : {account_data[0]}")
    print(f"Account Holder : {account_data[1]}")
    print(f"Current Balance: Rs. {float(account_data[4]):.2f}")
    print("=================================")


def main():
    initialize_excel()

    while True:
        print("\n==================================")
        print("      SIMPLE BANKING SYSTEM       ")
        print("==================================")
        print("1. Create Account")
        print("2. Deposit Money")
        print("3. Withdraw Money")
        print("4. Check Balance")
        print("5. Exit")
        print("==================================")

        choice = input("Enter your choice (1-5): ").strip()

        if choice == "1":
            create_account()
        elif choice == "2":
            deposit_money()
        elif choice == "3":
            withdraw_money()
        elif choice == "4":
            check_balance()
        elif choice == "5":
            print("\nThank you for using Simple Banking System! Have a great day!")
            break
        else:
            print("Error: Invalid Choice! Please select between 1 and 5.")


if __name__ == "__main__":
    main()